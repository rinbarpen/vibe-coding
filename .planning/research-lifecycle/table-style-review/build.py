from pathlib import Path
import json,csv,hashlib,zipfile
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).parent
T=[]
def add(key,title,heads,rows,widths,notes,groups=(),blocks=(),ref=None,style='journal_minimal'):
 T.append(dict(id=key,title=title,headers=heads,rows=rows,widths=widths,notes=notes,groups=groups,blocks=blocks,reference=ref,style=style))
add('01-main-results','Multi-dataset comparison',['Method','Accuracy (%)','F1 (%)','Accuracy (%)','F1 (%)'],[['Reference A','80.2','79.4','74.1','72.6'],['Reference B','82.1','81.0','75.8','74.3'],['Candidate','84.0','83.2','77.6','76.4']],[.32,.17,.17,.17,.17],['Invented fixture values; one run per configuration. Higher is better for every metric.','No best-value emphasis or significance claim. Dataset names are placeholders.'],groups=[(1,3,'Dataset A'),(3,5,'Dataset B')],style='grouped_comparison')
full=84.0
add('02-component-ablation','Component removal',['Configuration','A','B','C','Accuracy (%)','Delta (pp)'],[['Full','On','On','On','84.0','0.0'],['Without A','Off','On','On','81.2',f'{81.2-full:+.1f}'],['Without B','On','Off','On','82.0',f'{82-full:+.1f}'],['Without C','On','On','Off','83.1',f'{83.1-full:+.1f}']],[.34,.08,.08,.08,.22,.20],['Delta = accuracy minus Full accuracy, in percentage points (pp).','On/Off indicate component presence. Invented single-run fixture; no uncertainty estimates.'],ref=0,style='component_matrix')
vals=[78,81.2,82.9,84]
add('03-additive-ablation','Sequential component addition',['Configuration','Added','Accuracy (%)','Step delta (pp)'],[[name,comp,f'{v:.1f}','NA' if i==0 else f'{v-vals[i-1]:+.1f}'] for i,(name,comp,v) in enumerate(zip(['Base','Base + A','Base + A + B','Full'],['None','A','B','C'],vals))],[.40,.16,.22,.22],['Step delta is relative to the preceding row; NA = not applicable.','Order-dependent increments are not independent component effects. Invented fixture values.'],style='reference_delta')
add('04-design-ablation','Alternative design choices',['Choice','Setting','Accuracy (%)','Delta (pp)'],[['Fusion','Sum','82.2','-1.8'],['','Concatenate','83.1','-0.9'],['','Gated [ref]','84.0','0.0'],['Loss','Base','82.8','-1.2'],['','Combined [ref]','84.0','0.0']],[.20,.38,.22,.20],['Each block varies only its named factor; all other settings remain fixed.','[ref] defines the reference within each block. Deltas in pp; invented single-run fixtures.'],blocks=[3],style='grouped_comparison')
add('05-interaction-ablation','Two-factor interaction',['Factor A','B off','B on','B effect (pp)'],[['Off','78.0','80.0','+2.0'],['On','81.0','84.0','+3.0']],[.34,.22,.22,.22],['Cells contain accuracy (%). A x B contrast = (84 - 81) - (80 - 78) = +1.0 pp.','Descriptive contrast only: a single fixture run gives no uncertainty or interaction significance.'],groups=[(1,3,'Factor B')],style='factorial_matrix')
add('06-efficiency','Performance and resource trade-offs',['Method','Accuracy (%)','Params (M)','Latency (ms)','Memory (GB)'],[['Small','81.0','12.0','4.2','1.8'],['Medium','83.0','25.0','7.5','2.7'],['Large','84.0','48.0','13.1','4.5'],['Variant','82.5','30.0','9.2','3.2']],[.28,.18,.18,.18,.18],['Invented measurements: no actual hardware benchmark was performed.','Real tables must specify device, batch size, precision, warm-up and timing protocol.'],groups=[(1,2,'Performance'),(2,5,'Resources')],style='tradeoff')
add('07-robustness','Robustness under increasing corruption',['Condition','Reference (%)','Candidate (%)','Candidate drop (pp)'],[['Clean','80.2','84.0','0.0'],['Mild','76.1','81.0','3.0'],['Moderate','68.4','75.2','8.8'],['Severe','54.3','62.5','21.5']],[.34,.22,.22,.22],['Drop = clean candidate accuracy minus corrupted candidate accuracy; smaller is better.','Same evaluation protocol assumed for this fixture. Conditions and values are invented.'],ref=0,style='reference_delta')
runs=np.array([[80.1,81.0,80.6,79.9,80.4],[82.2,83.0,82.7,82.0,82.6]])
add('08-multiseed','Across-seed performance summary',['Method','Seeds, n','Mean (%)','Sample SD (pp)','Range (%)'],[[n,'5',f'{r.mean():.2f}',f'{r.std(ddof=1):.2f}',f'{r.min():.1f}–{r.max():.1f}'] for n,r in zip(['Reference','Candidate'],runs)],[.30,.14,.18,.20,.18],['Explicit multi-seed example; the workflow default remains one seed. SD uses ddof = 1.','Fixture seed IDs: 11, 22, 33, 44, 55. Values are invented; SD is computed, not invented.'],style='uncertainty_first')
for page in range(2):
 rows=[]
 for k in range(page*16,(page+1)*16):rows.append([f'config-{k+1:03d}',str(42),f'{76+(k%9)*.7:.1f}', ['Complete','Not run','Failed','Not applicable'][k%4],str(180) if k%4==0 else 'NA'])
 # No result reported for missing, failed, or inapplicable runs.
 for row in rows:
  if row[3]!='Complete':row[2]='NA'
 add(f'09-appendix-{page+1}',f'Complete configuration ledger ({page+1}/2)',['Configuration','Seed','Accuracy (%)','Status','Test n'],rows,[.29,.12,.20,.26,.13],['Fixture ledger; NA indicates no metric is reported. Status distinguishes why it is absent.','Headers repeat across pages; configuration IDs remain stable.'],style='supplement_dense')
# A/B style comparison uses identical component data, never altered values.
b=T[1]
add('10-style-reference', 'Same ablation / reference-delta layout',['Configuration','Accuracy (%)','Delta vs Full (pp)'],[[r[0],r[4],r[5]] for r in b['rows']],[.50,.25,.25],['Same values as component-removal table; compact variant for a space-limited manuscript.','Component definitions belong in the caption when the switch matrix is omitted.'],ref=0,style='reference_delta')
add('11-style-review','Same ablation / review-only heat shading',b['headers'],b['rows'],b['widths'],['Same values as component-removal table; colour is an optional review aid, not a significance cue.','Shading scales accuracy from 80 to 85%. Use the black-and-white variant for journal submission.'],ref=0,style='review_heat')
plt.rcParams.update({'font.family':'Liberation Sans','font.size':7,'pdf.fonttype':42,'svg.fonttype':'none'})
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69);sec.left_margin=sec.right_margin=Inches(.55);doc.styles['Normal'].font.name='Liberation Sans';doc.styles['Normal'].font.size=Pt(9)
review=[]
def alignment(t,j):
 if j==0 or (t['id']=='04-design-ablation' and j==1) or (t['id'].startswith('09-') and j==3):return 'left'
 if (t['id'] in ['02-component-ablation','11-style-review'] and j in [1,2,3]) or (t['id']=='03-additive-ablation' and j==1):return 'center'
 return 'right'
def xpos(t,e,j):
 mode=alignment(t,j)
 return e[j]+.008 if mode=='left' else ((e[j]+e[j+1])/2 if mode=='center' else e[j+1]-.012)
with PdfPages(R/'table-style-catalog.pdf') as pdf:
 for num,t in enumerate(T):
  n=len(t['rows']);group=bool(t['groups']);H=34+5*n+5*len(t['notes'])+5*group+3*len(t['blocks']);f=plt.figure(figsize=(183/25.4,H/25.4));a=f.add_axes([.025,0,.95,1]);a.set(xlim=(0,1),ylim=(H,0));a.axis('off');e=np.r_[0,np.cumsum(t['widths'])]
  a.text(0,5,f"{t['id']}  |  {t['title']}",fontsize=8,weight='bold',va='top');a.text(0,11,t['style']+'  /  SYNTHETIC FIXTURE — NOT RESEARCH EVIDENCE',fontsize=5.5,color='#555555',va='top');top=18;a.plot([0,1],[top,top],color='black',lw=.7)
  if group:
   for lo,hi,label in t['groups']:
    a.text((e[lo]+e[hi])/2,top+3,label,ha='center',va='center',fontsize=7);a.plot([e[lo]+.01,e[hi]-.01],[top+5.5]*2,color='black',lw=.4)
   top+=6
  for j,h in enumerate(t['headers']):a.text(xpos(t,e,j),top+3.5,h,ha=alignment(t,j),va='center',weight='bold',fontsize=7)
  yy=top+7;a.plot([0,1],[yy,yy],color='black',lw=.4); extents=[]
  for i,row in enumerate(t['rows']):
   if i in t['blocks']:yy+=3
   if t['style']=='review_heat':
    v=float(row[4]);a.add_patch(plt.Rectangle((e[4],yy),t['widths'][4],5,color=plt.cm.Blues(.08+.3*(v-80)/5),lw=0))
   for j,v in enumerate(row):
    txt=a.text(xpos(t,e,j),yy+2.7,v,ha=alignment(t,j),va='center',fontsize=7,weight='bold' if j==0 and i==t['reference'] else 'normal');extents.append((txt,j))
   yy+=5
  a.plot([0,1],[yy,yy],color='black',lw=.7)
  for i,note in enumerate(t['notes']):a.text(0,yy+4+i*4,note,fontsize=6,va='top')
  if t['reference'] is not None:a.text(0,yy+4+len(t['notes'])*4,'Bold row label identifies the reference, not statistical significance.',fontsize=6,va='top')
  f.canvas.draw();rend=f.canvas.get_renderer();over=[]
  for txt,j in extents:
   box=txt.get_window_extent(rend);low=a.transData.transform((e[j],0))[0];high=a.transData.transform((e[j+1],0))[0]
   if box.x0<low-1 or box.x1>high+1:over.append(txt.get_text())
  assert not over,(t['id'],over)
  for ext in ['png','svg','pdf']:f.savefig(R/(t['id']+'.'+ext),dpi=220)
  pdf.savefig(f);plt.close(f)
  with (R/(t['id']+'.csv')).open('w') as q:w=csv.writer(q);w.writerow(t['headers']);w.writerows(t['rows'])
  if num:doc.add_page_break()
  doc.add_paragraph(t['title'],'Heading 2');doc.add_paragraph('SYNTHETIC FIXTURE — NOT RESEARCH EVIDENCE\nStyle: '+t['style'])
  table=doc.add_table(rows=1+int(group),cols=len(t['headers']));table.autofit=False
  for j,ww in enumerate(t['widths']):table.columns[j].width=Inches(7.17*ww)
  if group:
   for lo,hi,label in t['groups']:
    cell=table.rows[0].cells[lo]
    if hi-lo>1:cell=cell.merge(table.rows[0].cells[hi-1])
    cell.text=label
  for j,h in enumerate(t['headers']):table.rows[-1].cells[j].text=h
  for row in t['rows']:
   for c,v in zip(table.add_row().cells,row):c.text=v
  for i,row in enumerate(table.rows):
   if i<1+int(group):row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
   for j,c in enumerate(row.cells):
    p=c.paragraphs[0];p.alignment={'left':WD_ALIGN_PARAGRAPH.LEFT,'right':WD_ALIGN_PARAGRAPH.RIGHT,'center':WD_ALIGN_PARAGRAPH.CENTER}[alignment(t,j)];p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
    for run in p.runs:run.font.size=Pt(8);run.bold=i<1+int(group) or (j==0 and i==1+int(group)+(t['reference'] if t['reference'] is not None else -99))
    borders=OxmlElement('w:tcBorders')
    for side in ['top','left','bottom','right']:
     z=OxmlElement('w:'+side);show=(side=='top' and i==0) or (side=='bottom' and i in [int(group),len(table.rows)-1]);z.set(qn('w:val'),'single' if show else 'nil');z.set(qn('w:sz'),'4');borders.append(z)
    c._tc.get_or_add_tcPr().append(borders)
  for note in t['notes']:doc.add_paragraph(note)
  review.append({'id':t['id'],'body_cell_overflows':over,'row_count':len(t['rows']),'fixture':True})
doc.save(R/'table-style-catalog.docx');re=Document(R/'table-style-catalog.docx');assert len(re.tables)==len(T)
for tt,t in zip(re.tables,T):assert [[c.text for c in row.cells] for row in list(tt.rows)[1+bool(t['groups']):]]==t['rows']
import statistics
assert all(np.isclose(row.std(ddof=1),statistics.stdev(row.tolist())) for row in runs)
(R/'fixtures.json').write_text(json.dumps({'tables':T,'seed_fixture':runs.tolist(),'seed_ids':[11,22,33,44,55]},indent=2));(R/'review.json').write_text(json.dumps({'tables':review,'native_word_tables_reopened':len(T),'multiseed_sd_verified':True,'word_gui_review':False},indent=2))
(R/'README.md').write_text('''# Table type × style catalog
All values are invented test fixtures, not executed experiments. Derived deltas and sample SD are computed from those fixtures. Default is one seed; only the explicitly labeled multi-seed page uses five.

PDF/PNG/SVG demonstrate layout. DOCX contains native editable tables; the review heat shading is preview-only and Word retains monochrome output. CSV and JSON preserve values. Word GUI pagination has not been visually tested.

Source: build.py. Main-result groups, component/removal/addition/design/interaction ablations, efficiency, robustness, multiseed, and appendix continuation are separate layouts. Pages 02, 10, 11 reuse identical removal data for style comparison. Bold row labels denote references, not best/significant results. No significance tests are claimed.

Nature-inspired monochrome editorial direction, not a universal Nature-family submission certification. No existing figure assets were changed.
''')
print('PASS:',len(T),'table pages; native Word values reopened; cell widths and sample SD verified.')
